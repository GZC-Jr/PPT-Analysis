import os
from pptx import Presentation
from PIL import Image
import io


# 注意：python-pptx本身不直接支持将幻灯片导出为图像。
# 最可靠的方法是使用外部COM自动化（仅限Windows）或调用libreoffice等工具。
# 这是一个巨大的依赖，并且不跨平台。
#
# 下面的实现是一个跨平台的“尽力而为”的替代方案，它提取幻灯片中的所有图像。
# 这不能完美复现幻灯片布局，但可以提供一些视觉内容。
# 如果需要完美渲染，建议使用`win32com`的实现（仅限Windows）。

def extract_slides_as_images(ppt_path, output_dir):
    """
    尽力而为的跨平台幻灯片图片提取方法。
    它会为每张幻灯片创建一个空白背景，然后将该幻灯片上的所有图片粘贴上去。
    这不能完美复现幻灯片，但提供了一种不依赖COM的解决方案。

    Args:
        ppt_path (str): PPTX文件路径.
        output_dir (str): 保存输出图片的目录.

    Returns:
        list: 生成的图片文件路径列表.
    """
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)

    image_paths = []
    try:
        prs = Presentation(ppt_path)

        # 获取演示文稿的尺寸（以像素为单位，假设96 DPI）
        width = int(prs.slide_width.pt * 96 / 72)
        height = int(prs.slide_height.pt * 96 / 72)

        for i, slide in enumerate(prs.slides):
            # 为每张幻灯片创建一个白色背景的画布
            slide_image = Image.new('RGB', (width, height), 'white')

            # 遍历幻灯片中的所有形状，寻找图片
            for shape in slide.shapes:
                if hasattr(shape, "image"):
                    image = shape.image
                    # 获取图片数据并用Pillow打开
                    image_bytes = image.blob
                    img = Image.open(io.BytesIO(image_bytes))

                    # 获取图片在幻灯片上的位置和大小
                    left = int(shape.left.pt * 96 / 72)
                    top = int(shape.top.pt * 96 / 72)

                    # 调整图片大小并粘贴到画布上
                    if hasattr(shape, "width") and hasattr(shape, "height"):
                        shape_w = int(shape.width.pt * 96 / 72)
                        shape_h = int(shape.height.pt * 96 / 72)
                        img = img.resize((shape_w, shape_h), Image.LANCZOS)

                    slide_image.paste(img, (left, top), mask=img.split()[3] if img.mode == 'RGBA' else None)

            # 保存最终合成的图片
            image_path = os.path.join(output_dir, f"slide_{i + 1}.png")
            slide_image.save(image_path)
            image_paths.append(image_path)

    except Exception as e:
        print(f"提取幻灯片时出错: {e}")
        # 如果出错，可以返回一些占位符
        return []

    return image_paths

# --- 备选方案 (仅限Windows，需要 pywin32) ---
# def extract_slides_as_images_win32(ppt_path, output_dir):
#     import win32com.client
#     import pythoncom
#
#     pythoncom.CoInitialize()
#     powerpoint = win32com.client.Dispatch("Powerpoint.Application")
#     powerpoint.Visible = 1
#
#     if not os.path.exists(output_dir):
#         os.makedirs(output_dir)
#
#     image_paths = []
#     try:
#         presentation = powerpoint.Presentations.Open(os.path.abspath(ppt_path), WithWindow=False)
#         for i, slide in enumerate(presentation.Slides):
#             image_path = os.path.join(output_dir, f"slide_{i+1}.png")
#             slide.Export(os.path.abspath(image_path), "PNG")
#             image_paths.append(image_path)
#         presentation.Close()
#     except Exception as e:
#         print(f"Error extracting slides with win32com: {e}")
#     finally:
#         powerpoint.Quit()
#         pythoncom.CoUninitialize()
#
#     return image_paths

